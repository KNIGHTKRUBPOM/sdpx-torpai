from __future__ import annotations

import json
import os
import csv
import io
import re
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from typing import Annotated, Literal
from uuid import uuid4
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import anyio.to_thread
from fastapi import Depends, FastAPI, File, Header, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from fastapi.responses import Response
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from src.domain.models import PairAssignment, Student
from src.persistence.database import Base, SessionLocal, engine, get_session
from src.persistence.migrations import migrate_existing_schema
from src.persistence.models import (
    AppealRecord,
    AssignmentRecord,
    AuditRecord,
    ClassroomRecord,
    ComparisonExclusionRecord,
    ComparisonRecord,
    CriterionRecord,
    GroupRecord,
    MembershipRecord,
    NotificationRecord,
    PairAssignmentRecord,
    PrivacyAcknowledgementRecord,
    ScoreOverrideRecord,
    ScoreSnapshotRecord,
    SubmissionRevisionRecord,
    UserRecord,
)
from src.persistence.seed import seed_demo_data, stable_id
from src.repositories.sqlalchemy_pair_assignment_repository import SqlAlchemyPairAssignmentRepository
from src.services.pairing_service import PairingNotFeasibleError, PairingService
from src.services.roster_service import RosterValidationError, normalize_email, parse_roster_csv
from src.services.score_reporting_service import ScoreReportingService
from src.services.quality_signal_service import QualitySignalService
from src.services.maintenance_service import finalize_scores, run_maintenance


AUTH_MODE = os.getenv("AUTH_MODE", "demo")
ALLOWED_ORIGINS = [
    value.strip()
    for value in os.getenv("ALLOWED_ORIGINS", "http://localhost:5173,http://localhost:8080").split(",")
    if value.strip()
]


class ApiError(Exception):
    def __init__(self, status_code: int, code: str, message: str, field: str | None = None) -> None:
        self.status_code = status_code
        self.code = code
        self.message = message
        self.field = field


class ApiModel(BaseModel):
    model_config = ConfigDict(populate_by_name=True)


class ClassroomCreate(ApiModel):
    name: str = Field(min_length=2, max_length=160)
    slug: str = Field(min_length=2, max_length=120, pattern=r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
    timezone: str = Field(default="Asia/Bangkok", min_length=1, max_length=80)

    @field_validator("timezone")
    @classmethod
    def require_iana_timezone(cls, value: str) -> str:
        try:
            ZoneInfo(value)
        except ZoneInfoNotFoundError as error:
            raise ValueError("timezone must be a valid IANA timezone") from error
        return value


class CriterionCreate(ApiModel):
    name: str = Field(min_length=2, max_length=160)
    prompt: str = Field(min_length=5, max_length=2000)
    weight_pct: Decimal = Field(alias="weightPct", gt=0, le=100)
    side: Literal["GROUP", "INDIVIDUAL"]


class AssignmentCreate(ApiModel):
    name: str = Field(min_length=2, max_length=180)
    description: str = Field(default="", max_length=4000)
    deadline: datetime
    individual_deadline: datetime | None = Field(default=None, alias="individualDeadline")
    group_max_score: Decimal = Field(default=Decimal("15"), alias="groupMaxScore", ge=0)
    individual_max_score: Decimal = Field(default=Decimal("0"), alias="individualMaxScore", ge=0)
    criterion_name: str | None = Field(default=None, alias="criterionName", min_length=2, max_length=160)
    criterion_prompt: str | None = Field(default=None, alias="criterionPrompt", min_length=5, max_length=2000)
    criteria: list[CriterionCreate] = Field(default_factory=list)
    target_coverage: int = Field(default=5, alias="targetCoverage", ge=1, le=20)
    max_workload: int = Field(default=8, alias="maxWorkload", ge=1, le=30)
    min_comparisons: int = Field(default=3, alias="minComparisons", ge=1, le=20)
    score_floor: Decimal = Field(default=Decimal("0.60"), alias="scoreFloor", ge=0, le=1)
    score_ceiling: Decimal = Field(default=Decimal("1.00"), alias="scoreCeiling", ge=0, le=1)
    completion_threshold: Decimal = Field(default=Decimal("0.90"), alias="completionThreshold", gt=0, le=1)
    instructor_weight: Decimal = Field(default=Decimal("1.00"), alias="instructorWeight", ge=0)
    seed: int = 20260819

    @field_validator("deadline", "individual_deadline")
    @classmethod
    def require_timezone(cls, value: datetime | None) -> datetime | None:
        if value is None:
            return None
        if value.tzinfo is None:
            raise ValueError("deadline must include a timezone offset")
        return value.astimezone(timezone.utc)

    @model_validator(mode="after")
    def validate_scoring_configuration(self) -> "AssignmentCreate":
        if self.group_max_score + self.individual_max_score <= 0:
            raise ValueError("at least one evaluation side must have a positive maximum score")
        if self.score_floor >= self.score_ceiling:
            raise ValueError("scoreFloor must be lower than scoreCeiling")
        if self.criteria:
            for side, maximum in (("GROUP", self.group_max_score), ("INDIVIDUAL", self.individual_max_score)):
                side_criteria = [criterion for criterion in self.criteria if criterion.side == side]
                if maximum > 0 and not side_criteria:
                    raise ValueError(f"{side} criteria are required when its maximum score is positive")
                total = sum((criterion.weight_pct for criterion in side_criteria), start=Decimal("0"))
                if side_criteria and abs(total - Decimal("100")) > Decimal("0.01"):
                    raise ValueError(f"{side} criterion weights must total 100%")
        elif not self.criterion_name or not self.criterion_prompt:
            raise ValueError("criteria or criterionName/criterionPrompt are required")
        return self
class PublishRequest(ApiModel):
    seed: int | None = None


class ComparisonSave(ApiModel):
    choice: int = Field(ge=1, le=6)
    time_on_task_ms: int = Field(default=0, alias="timeOnTaskMs", ge=0)


class SubmissionRequest(ApiModel):
    side: Literal["GROUP", "INDIVIDUAL"] = "GROUP"


class ExportRequest(ApiModel):
    format: Literal["CSV", "XLSX"] = "CSV"
    report: Literal["GROUP", "INDIVIDUAL", "COVERAGE", "RAW"]
    include_identities: bool = Field(default=False, alias="includeIdentities")
    identity_confirmation: Literal["EXPORT_IDENTITIES"] | None = Field(default=None, alias="identityConfirmation")


class FinalizeRequest(ApiModel):
    reason: str | None = Field(default=None, max_length=2000)


class ReasonRequest(ApiModel):
    reason: str = Field(min_length=5, max_length=2000)


class ScoreOverrideCreate(ReasonRequest):
    side: Literal["GROUP", "INDIVIDUAL"]
    item_id: str = Field(alias="itemId", min_length=1, max_length=64)
    override_value: Decimal = Field(alias="overrideValue", ge=0)


class AppealCreate(ApiModel):
    assignment_id: str = Field(alias="assignmentId", min_length=1, max_length=64)
    message: str = Field(min_length=10, max_length=4000)


class AppealResolve(ApiModel):
    status: Literal["RESOLVED", "REJECTED"]
    resolution: str = Field(min_length=5, max_length=4000)


class PrivacyAcknowledgement(ApiModel):
    notice_version: Literal["v1.0"] = Field(default="v1.0", alias="noticeVersion")


class NotificationRead(ApiModel):
    read: bool = True


SessionDep = Annotated[Session, Depends(get_session, scope="function")]


def initialize_database() -> None:
    Base.metadata.create_all(bind=engine)
    migrate_existing_schema(engine)
    with SessionLocal() as session:
        if AUTH_MODE == "demo":
            seed_demo_data(session)


if os.getenv("SKIP_DATABASE_INITIALIZATION") != "1":
    initialize_database()


@asynccontextmanager
async def application_lifespan(_: FastAPI):
    # A synchronous SQLAlchemy session holds a connection for the request.
    # Do not let more request threads enter than the per-worker connection pool
    # can serve, otherwise cleanup can be starved during a concurrent spike.
    anyio.to_thread.current_default_thread_limiter().total_tokens = int(
        os.getenv("THREAD_POOL_SIZE", os.getenv("DB_POOL_SIZE", "15"))
    )
    yield


app = FastAPI(
    title="PairEval API",
    version="2.0.0-m4-local",
    description="PairEval local M4 API for evaluation, explainable scoring, governance, notifications and exports.",
    lifespan=application_lifespan,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def attach_request_id(request: Request, call_next):
    request_id = request.headers.get("X-Request-ID", str(uuid4()))
    request.state.request_id = request_id
    started = datetime.now(timezone.utc)
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    duration_ms = int((datetime.now(timezone.utc) - started).total_seconds() * 1000)
    print(
        json.dumps(
            {
                "event": "http_request",
                "requestId": request_id,
                "method": request.method,
                "path": request.url.path,
                "status": response.status_code,
                "duration_ms": duration_ms,
            },
            ensure_ascii=False,
        )
    )
    return response


@app.exception_handler(ApiError)
async def handle_api_error(request: Request, error: ApiError) -> JSONResponse:
    return JSONResponse(
        status_code=error.status_code,
        content={
            "error": {
                "code": error.code,
                "message": error.message,
                "field": error.field,
                "requestId": request.state.request_id,
            }
        },
    )


@app.exception_handler(RosterValidationError)
async def handle_roster_error(request: Request, error: RosterValidationError) -> JSONResponse:
    return JSONResponse(
        status_code=400,
        content={
            "error": {
                "code": "ROSTER_INVALID",
                "message": "Roster CSV was rejected; no rows were imported.",
                "details": error.errors,
                "requestId": request.state.request_id,
            }
        },
    )


@app.exception_handler(RequestValidationError)
async def handle_request_validation_error(request: Request, error: RequestValidationError) -> JSONResponse:
    details = [
        {
            "field": ".".join(str(part) for part in item["loc"] if part not in {"body", "query", "path", "header"}),
            "message": item["msg"],
            "type": item["type"],
        }
        for item in error.errors()
    ]
    return JSONResponse(
        status_code=422,
        content={
            "error": {
                "code": "REQUEST_VALIDATION_FAILED",
                "message": "The request contains invalid data.",
                "details": details,
                "requestId": request.state.request_id,
            }
        },
    )


def current_user(
    session: SessionDep,
    user_id: Annotated[str | None, Header(alias="X-User-Id")] = None,
) -> UserRecord:
    if not user_id:
        raise ApiError(401, "AUTH_REQUIRED", "Choose a demo account before continuing.")
    user = session.get(UserRecord, user_id)
    if user is None:
        raise ApiError(401, "AUTH_USER_UNKNOWN", "The authenticated user does not exist.")
    return user


UserDep = Annotated[UserRecord, Depends(current_user)]


def request_id(request: Request) -> str:
    return str(request.state.request_id)


def classroom_membership(
    session: Session,
    classroom_id: str,
    user: UserRecord,
    allowed_roles: set[str] | None = None,
) -> tuple[ClassroomRecord, MembershipRecord]:
    row = session.execute(
        select(ClassroomRecord, MembershipRecord)
        .join(MembershipRecord, MembershipRecord.classroom_id == ClassroomRecord.id)
        .where(
            ClassroomRecord.id == classroom_id,
            MembershipRecord.classroom_id == classroom_id,
            MembershipRecord.user_id == user.id,
        )
    ).one_or_none()
    if row is None:
        raise ApiError(404, "CLASSROOM_NOT_FOUND", "Classroom was not found.")
    classroom, membership = row
    if allowed_roles is not None and membership.role not in allowed_roles:
        raise ApiError(403, "CLASSROOM_FORBIDDEN", "Your role cannot perform this action.")
    return classroom, membership


def assignment_access(
    session: Session,
    assignment_id: str,
    user: UserRecord,
    allowed_roles: set[str] | None = None,
) -> tuple[AssignmentRecord, ClassroomRecord, MembershipRecord]:
    row = session.execute(
        select(AssignmentRecord, ClassroomRecord, MembershipRecord)
        .join(ClassroomRecord, ClassroomRecord.id == AssignmentRecord.classroom_id)
        .join(MembershipRecord, MembershipRecord.classroom_id == ClassroomRecord.id)
        .where(
            AssignmentRecord.id == assignment_id,
            MembershipRecord.user_id == user.id,
        )
    ).one_or_none()
    if row is None:
        raise ApiError(404, "ASSIGNMENT_NOT_FOUND", "Assignment was not found.")
    assignment, classroom, membership = row
    if allowed_roles is not None and membership.role not in allowed_roles:
        raise ApiError(403, "CLASSROOM_FORBIDDEN", "Your role cannot perform this action.")
    return assignment, classroom, membership


def aware_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def evaluation_deadline(assignment: AssignmentRecord, side: str) -> datetime:
    value = assignment.individual_deadline if side == "INDIVIDUAL" else assignment.deadline
    if value is None:
        return aware_utc(assignment.deadline)
    return aware_utc(value)


def audit(
    session: Session,
    actor: UserRecord,
    action: str,
    resource_type: str,
    resource_id: str,
    req_id: str,
    details: dict[str, object] | None = None,
    before: dict[str, object] | None = None,
    after: dict[str, object] | None = None,
    reason: str | None = None,
    ip_address: str | None = None,
    assignment_id: str | None = None,
) -> None:
    session.add(
        AuditRecord(
            id=str(uuid4()),
            actor_user_id=actor.id,
            assignment_id=assignment_id or (resource_id if resource_type == "assignment" else None),
            action=action,
            resource_type=resource_type,
            resource_id=resource_id,
            details_json=json.dumps(details or {}, ensure_ascii=False, sort_keys=True),
            before_json=json.dumps(before or {}, ensure_ascii=False, sort_keys=True),
            after_json=json.dumps(after or {}, ensure_ascii=False, sort_keys=True),
            reason=reason,
            ip_address=ip_address,
            request_id=req_id,
        )
    )


def queue_assignment_notifications(
    session: Session,
    assignment: AssignmentRecord,
    notification_type: str,
    message: str,
    user_ids: list[str] | None = None,
) -> int:
    if user_ids is None:
        user_ids = list(
            session.scalars(
                select(MembershipRecord.user_id).where(
                    MembershipRecord.classroom_id == assignment.classroom_id,
                    MembershipRecord.role == "STUDENT",
                    MembershipRecord.status == "ACTIVE",
                )
            ).all()
        )
    queued = 0
    for user_id in user_ids:
        dedupe_key = f"{notification_type}:{assignment.id}:{user_id}"
        if session.scalar(select(NotificationRecord.id).where(NotificationRecord.dedupe_key == dedupe_key)):
            continue
        session.add(
            NotificationRecord(
                id=str(uuid4()),
                dedupe_key=dedupe_key,
                user_id=user_id,
                notification_type=notification_type,
                payload_json=json.dumps(
                    {
                        "assignmentId": assignment.id,
                        "message": message,
                        "path": f"/assignments/{assignment.id}",
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                ),
            )
        )
        queued += 1
    return queued


def ensure_instructor(session: Session, user: UserRecord) -> None:
    role = session.scalar(
        select(MembershipRecord.role).where(
            MembershipRecord.user_id == user.id,
            MembershipRecord.role.in_(["OWNER", "INSTRUCTOR"]),
        )
    )
    if role is None:
        raise ApiError(403, "INSTRUCTOR_REQUIRED", "An instructor account is required.")


def assignment_payload(session: Session, assignment: AssignmentRecord, user_id: str) -> dict[str, object]:
    classroom = session.get(ClassroomRecord, assignment.classroom_id)
    criteria = session.scalars(
        select(CriterionRecord).where(CriterionRecord.assignment_id == assignment.id).order_by(CriterionRecord.id)
    ).all()
    progress_rows = session.execute(
        select(
            CriterionRecord.side,
            func.count(PairAssignmentRecord.id),
            func.count(ComparisonRecord.pair_assignment_id),
        )
        .join(PairAssignmentRecord, PairAssignmentRecord.criterion_id == CriterionRecord.id)
        .outerjoin(ComparisonRecord, ComparisonRecord.pair_assignment_id == PairAssignmentRecord.id)
        .where(
            PairAssignmentRecord.assignment_id == assignment.id,
            PairAssignmentRecord.evaluator_id == user_id,
        )
        .group_by(CriterionRecord.side)
    ).all()
    assigned_by_side = {side: assigned_count for side, assigned_count, _ in progress_rows}
    answered_by_side = {side: answered_count for side, _, answered_count in progress_rows}
    assigned = sum(assigned_by_side.values())
    answered = sum(answered_by_side.values())
    return {
        "id": assignment.id,
        "classroomId": assignment.classroom_id,
        "classroomName": classroom.name if classroom else "",
        "name": assignment.name,
        "description": assignment.description,
        "status": assignment.status,
        "deadline": aware_utc(assignment.deadline).isoformat(),
        "individualDeadline": aware_utc(assignment.individual_deadline).isoformat() if assignment.individual_deadline else None,
        "groupMaxScore": float(assignment.group_max_score),
        "individualMaxScore": float(assignment.individual_max_score),
        "scoreFloor": float(assignment.score_floor),
        "scoreCeiling": float(assignment.score_ceiling),
        "completionThreshold": float(assignment.completion_threshold),
        "instructorWeight": float(assignment.instructor_weight),
        "minComparisons": assignment.min_comparisons,
        "criteria": [
            {
                "id": criterion.id,
                "name": criterion.name,
                "prompt": criterion.prompt,
                "weightPct": float(criterion.weight_pct),
                "side": criterion.side,
            }
            for criterion in criteria
        ],
        "assignedCount": assigned,
        "answeredCount": answered,
        "groupAssignedCount": assigned_by_side.get("GROUP", 0),
        "groupAnsweredCount": answered_by_side.get("GROUP", 0),
        "individualAssignedCount": assigned_by_side.get("INDIVIDUAL", 0),
        "individualAnsweredCount": answered_by_side.get("INDIVIDUAL", 0),
    }


@app.get("/")
def root():
    return {"service": "PairEval API", "version": app.version, "status": "ok", "mode": AUTH_MODE}


@app.get("/health")
def health(session: SessionDep):
    session.execute(select(1))
    return {"status": "ok", "database": "connected"}


@app.get("/api/demo/users")
def list_demo_users(session: SessionDep):
    if AUTH_MODE != "demo":
        raise ApiError(404, "DEMO_DISABLED", "Demo authentication is disabled.")
    rows = session.execute(
        select(UserRecord, MembershipRecord.role)
        .join(MembershipRecord, MembershipRecord.user_id == UserRecord.id)
        .order_by(MembershipRecord.role, UserRecord.id)
    ).all()
    role_rank = {"OWNER": 3, "INSTRUCTOR": 2, "STUDENT": 1}
    users: dict[str, dict[str, str]] = {}
    for user, role in rows:
        current = users.get(user.id)
        if current is None or role_rank[role] > role_rank[current["role"]]:
            users[user.id] = {
                "id": user.id,
                "email": user.email,
                "displayName": user.display_name,
                "role": role,
            }
    return sorted(users.values(), key=lambda item: (-role_rank[item["role"]], item["id"]))


@app.get("/api/me")
def get_me(session: SessionDep, user: UserDep):
    memberships = session.execute(
        select(MembershipRecord, ClassroomRecord)
        .join(ClassroomRecord, ClassroomRecord.id == MembershipRecord.classroom_id)
        .where(MembershipRecord.user_id == user.id)
    ).all()
    return {
        "id": user.id,
        "email": user.email,
        "displayName": user.display_name,
        "privacyAcknowledged": session.get(PrivacyAcknowledgementRecord, user.id) is not None,
        "memberships": [
            {
                "classroomId": membership.classroom_id,
                "classroomName": classroom.name,
                "role": membership.role,
                "groupId": membership.group_id,
            }
            for membership, classroom in memberships
        ],
    }


@app.get("/api/classrooms")
def list_classrooms(session: SessionDep, user: UserDep):
    rows = session.execute(
        select(ClassroomRecord, MembershipRecord)
        .join(MembershipRecord, MembershipRecord.classroom_id == ClassroomRecord.id)
        .where(MembershipRecord.user_id == user.id)
        .order_by(ClassroomRecord.created_at.desc())
    ).all()
    return [
        {
            "id": classroom.id,
            "name": classroom.name,
            "slug": classroom.slug,
            "timezone": classroom.timezone,
            "role": membership.role,
        }
        for classroom, membership in rows
    ]


@app.post("/api/classrooms", status_code=201)
def create_classroom(payload: ClassroomCreate, request: Request, session: SessionDep, user: UserDep):
    ensure_instructor(session, user)
    classroom = ClassroomRecord(
        id=str(uuid4()),
        name=payload.name.strip(),
        slug=payload.slug,
        timezone=payload.timezone,
        owner_user_id=user.id,
    )
    session.add(classroom)
    session.add(MembershipRecord(id=str(uuid4()), classroom_id=classroom.id, user_id=user.id, role="OWNER"))
    audit(session, user, "CLASSROOM_CREATED", "classroom", classroom.id, request_id(request))
    try:
        session.commit()
    except IntegrityError as error:
        session.rollback()
        raise ApiError(409, "CLASSROOM_SLUG_EXISTS", "This classroom slug is already in use.", "slug") from error
    return {"id": classroom.id, "name": classroom.name, "slug": classroom.slug, "timezone": classroom.timezone}


@app.post("/api/classrooms/{classroom_id}/roster:import", status_code=201)
async def import_roster(
    classroom_id: str,
    request: Request,
    session: SessionDep,
    user: UserDep,
    file: UploadFile = File(...),
):
    classroom, _ = classroom_membership(session, classroom_id, user, {"OWNER", "INSTRUCTOR"})
    raw = await file.read(1_000_001)
    if len(raw) > 1_000_000:
        raise ApiError(413, "ROSTER_TOO_LARGE", "Roster CSV must not exceed 1 MB.", "file")
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ApiError(400, "ROSTER_ENCODING", "Roster CSV must use UTF-8.", "file") from error
    rows = parse_roster_csv(content)
    published = session.scalar(
        select(func.count()).select_from(AssignmentRecord).where(
            AssignmentRecord.classroom_id == classroom_id,
            AssignmentRecord.status != "DRAFT",
        )
    ) or 0
    if published:
        raise ApiError(409, "ROSTER_LOCKED", "Unpublish assignments before replacing this roster.")

    session.execute(delete(MembershipRecord).where(MembershipRecord.classroom_id == classroom_id, MembershipRecord.role == "STUDENT"))
    session.execute(delete(GroupRecord).where(GroupRecord.classroom_id == classroom_id))
    session.flush()
    groups: dict[str, GroupRecord] = {}
    for row in rows:
        if row.group_name not in groups:
            group = GroupRecord(
                id=stable_id("group", f"{classroom_id}:{row.group_name}"),
                classroom_id=classroom_id,
                name=row.group_name,
                artifact_url=row.artifact_url,
            )
            groups[row.group_name] = group
            session.add(group)
        elif row.artifact_url and not groups[row.group_name].artifact_url:
            groups[row.group_name].artifact_url = row.artifact_url
    session.flush()
    for row in rows:
        existing = session.scalar(select(UserRecord).where(UserRecord.email == row.email))
        if existing is None:
            existing = UserRecord(
                id=stable_id("user", row.email),
                email=normalize_email(row.email),
                display_name=row.display_name,
                status="PENDING",
            )
            session.add(existing)
            session.flush()
        session.add(
            MembershipRecord(
                id=stable_id("membership", f"{classroom_id}:{existing.id}"),
                classroom_id=classroom_id,
                user_id=existing.id,
                role="STUDENT",
                group_id=groups[row.group_name].id,
                student_id=row.student_id,
            )
        )
    audit(
        session,
        user,
        "ROSTER_IMPORTED",
        "classroom",
        classroom.id,
        request_id(request),
        {"students": len(rows), "groups": len(groups)},
    )
    session.commit()
    return {"classroomId": classroom.id, "studentCount": len(rows), "groupCount": len(groups), "atomic": True}


@app.post("/api/classrooms/{classroom_id}/assignments", status_code=201)
def create_assignment(
    classroom_id: str,
    payload: AssignmentCreate,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    classroom, _ = classroom_membership(session, classroom_id, user, {"OWNER", "INSTRUCTOR"})
    if payload.deadline <= datetime.now(timezone.utc):
        raise ApiError(422, "DEADLINE_IN_PAST", "Deadline must be in the future.", "deadline")
    assignment = AssignmentRecord(
        id=str(uuid4()),
        classroom_id=classroom.id,
        name=payload.name.strip(),
        description=payload.description.strip(),
        status="DRAFT",
        group_max_score=payload.group_max_score,
        individual_max_score=payload.individual_max_score,
        deadline=payload.deadline,
        individual_deadline=(payload.individual_deadline or payload.deadline) if payload.individual_max_score > 0 else None,
        target_coverage=payload.target_coverage,
        max_workload=payload.max_workload,
        min_comparisons=payload.min_comparisons,
        score_floor=payload.score_floor,
        score_ceiling=payload.score_ceiling,
        completion_threshold=payload.completion_threshold,
        instructor_weight=payload.instructor_weight,
        scoring_formula_version="v2.0",
        seed=payload.seed,
    )
    criterion_inputs = payload.criteria or [
        CriterionCreate(
            name=payload.criterion_name or "Group quality",
            prompt=payload.criterion_prompt or "ผลงานใดมีคุณภาพดีกว่า?",
            weightPct=Decimal("100"),
            side="GROUP",
        )
    ]
    criteria = [
        CriterionRecord(
            id=str(uuid4()),
            assignment_id=assignment.id,
            name=item.name.strip(),
            prompt=item.prompt.strip(),
            weight_pct=item.weight_pct,
            side=item.side,
        )
        for item in criterion_inputs
    ]
    session.add_all([assignment, *criteria])
    audit(session, user, "ASSIGNMENT_CREATED", "assignment", assignment.id, request_id(request))
    session.commit()
    return assignment_payload(session, assignment, user.id)


@app.get("/api/assignments")
def list_assignments(session: SessionDep, user: UserDep):
    memberships = session.scalars(select(MembershipRecord).where(MembershipRecord.user_id == user.id)).all()
    classroom_ids = [membership.classroom_id for membership in memberships]
    if not classroom_ids:
        return []
    instructor_ids = {
        membership.classroom_id for membership in memberships if membership.role in {"OWNER", "INSTRUCTOR"}
    }
    assignments = session.scalars(
        select(AssignmentRecord)
        .where(AssignmentRecord.classroom_id.in_(classroom_ids))
        .order_by(AssignmentRecord.created_at.desc())
    ).all()
    return [
        assignment_payload(session, item, user.id)
        for item in assignments
        if item.classroom_id in instructor_ids or item.status in {"PUBLISHED", "CLOSED", "FINALIZED"}
    ]


def feasibility_for_assignment(session: Session, assignment: AssignmentRecord):
    members = session.scalars(
        select(MembershipRecord).where(
            MembershipRecord.classroom_id == assignment.classroom_id,
            MembershipRecord.role == "STUDENT",
            MembershipRecord.status == "ACTIVE",
            MembershipRecord.group_id.is_not(None),
        )
    ).all()
    students = [Student(id=member.user_id, group_id=member.group_id or "") for member in members]
    groups = session.scalars(
        select(GroupRecord).where(GroupRecord.classroom_id == assignment.classroom_id).order_by(GroupRecord.name)
    ).all()
    if assignment.group_max_score <= 0:
        group_result = {
            "targetCoverage": assignment.target_coverage,
            "actualCoverage": 0,
            "workload": 0,
            "pairCount": 0,
            "totalComparisons": 0,
            "reduced": False,
            "explanation": "ปิดการประเมินแบบกลุ่มสำหรับ Assignment นี้",
        }
    elif len(groups) == 2:
        group_result = {
            "targetCoverage": assignment.target_coverage,
            "actualCoverage": 1,
            "workload": 0,
            "pairCount": 1,
            "totalComparisons": 1,
            "reduced": True,
            "explanation": (
                "ห้องนี้มี 2 กลุ่ม นักศึกษาไม่สามารถประเมินโดยไม่เห็นกลุ่มตัวเอง "
                "จึงสร้างคู่สำหรับผู้สอน 1 ครั้งต่อเกณฑ์"
            ),
        }
    else:
        try:
            result = PairingService(SqlAlchemyPairAssignmentRepository(session)).solve_group_feasibility(
                students,
                target_coverage=assignment.target_coverage,
                max_workload=assignment.max_workload,
            )
        except PairingNotFeasibleError as error:
            raise ApiError(409, "PAIRING_NOT_FEASIBLE", str(error)) from error
        group_result = {
            "targetCoverage": result.target_coverage,
            "actualCoverage": result.actual_coverage,
            "workload": result.workload,
            "pairCount": result.pair_count,
            "totalComparisons": result.total_comparisons,
            "reduced": result.reduced,
            "explanation": result.explanation,
        }
    individual = []
    members_by_group: dict[str, int] = {}
    for member in members:
        if member.group_id:
            members_by_group[member.group_id] = members_by_group.get(member.group_id, 0) + 1
    if assignment.individual_max_score > 0:
        for group in groups:
            result = PairingService.solve_individual_feasibility(
                members_by_group.get(group.id, 0),
                max_workload=assignment.max_workload,
            )
            individual.append(
                {
                    "groupId": group.id,
                    "groupName": group.name,
                    "groupSize": result.group_size,
                    "enabled": result.enabled,
                    "pairCount": result.pair_count,
                    "workload": result.workload,
                    "minCoverage": result.min_coverage,
                    "maxCoverage": result.max_coverage,
                    "lowConfidence": result.low_confidence,
                    "workloadCapped": result.workload_capped,
                }
            )
    return {**group_result, "group": group_result, "individual": individual}


@app.get("/api/assignments/{assignment_id}/feasibility")
def get_feasibility(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    return feasibility_for_assignment(session, assignment)


@app.post("/api/assignments/{assignment_id}:publish", status_code=201)
def publish_assignment(
    assignment_id: str,
    payload: PublishRequest,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    assignment, classroom, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    if assignment.status != "DRAFT":
        raise ApiError(409, "ASSIGNMENT_ALREADY_PUBLISHED", "Only a draft assignment can be published.")
    criteria = session.scalars(select(CriterionRecord).where(CriterionRecord.assignment_id == assignment.id)).all()
    for side, maximum in (("GROUP", assignment.group_max_score), ("INDIVIDUAL", assignment.individual_max_score)):
        side_criteria = [criterion for criterion in criteria if criterion.side == side]
        weight_total = sum((criterion.weight_pct for criterion in side_criteria), start=Decimal("0"))
        if maximum > 0 and (not side_criteria or abs(weight_total - Decimal("100")) > Decimal("0.01")):
            raise ApiError(422, "CRITERIA_WEIGHT_INVALID", f"{side.title()} criterion weights must total 100%.")
    groups = session.scalars(select(GroupRecord).where(GroupRecord.classroom_id == assignment.classroom_id)).all()
    if assignment.group_max_score > 0 and any(not group.artifact_url for group in groups):
        raise ApiError(422, "ARTIFACT_REQUIRED", "Every group must have an artifact_url before publish.")
    members = session.scalars(
        select(MembershipRecord).where(
            MembershipRecord.classroom_id == assignment.classroom_id,
            MembershipRecord.role == "STUDENT",
            MembershipRecord.status == "ACTIVE",
            MembershipRecord.group_id.is_not(None),
        )
    ).all()
    students = [Student(id=member.user_id, group_id=member.group_id or "") for member in members]
    service = PairingService(SqlAlchemyPairAssignmentRepository(session))
    pair_count = 0
    try:
        effective_seed = payload.seed if payload.seed is not None else assignment.seed
        group_criteria = [criterion for criterion in criteria if criterion.side == "GROUP"]
        individual_criteria = [criterion for criterion in criteria if criterion.side == "INDIVIDUAL"]
        if assignment.group_max_score > 0 and len(groups) == 2:
            item_a, item_b = sorted(group.id for group in groups)
            for criterion in group_criteria:
                pair = PairAssignment(
                    id=stable_id("two-group-pair", f"{assignment.id}:{criterion.id}"),
                    assignment_id=assignment.id,
                    criterion_id=criterion.id,
                    evaluator_id=classroom.owner_user_id,
                    item_a_id=item_a,
                    item_b_id=item_b,
                    display_left_item_id=(item_a, item_b)[effective_seed % 2],
                )
                service.repository.replace_for_criterion(assignment.id, criterion.id, [pair])
                pair_count += 1
        elif assignment.group_max_score > 0:
            for criterion in group_criteria:
                _, generated = service.generate_group_pairs(
                    assignment.id,
                    criterion.id,
                    students,
                    seed=effective_seed,
                    target_coverage=assignment.target_coverage,
                    max_workload=assignment.max_workload,
                )
                pair_count += len(generated)
        if assignment.individual_max_score > 0:
            for criterion in individual_criteria:
                _, generated = service.generate_individual_pairs(
                    assignment.id,
                    criterion.id,
                    students,
                    seed=effective_seed,
                    max_workload=assignment.max_workload,
                )
                pair_count += len(generated)
    except PairingNotFeasibleError as error:
        session.rollback()
        raise ApiError(409, "PAIRING_NOT_FEASIBLE", str(error)) from error
    assignment.status = "PUBLISHED"
    assignment.published_at = datetime.now(timezone.utc)
    if payload.seed is not None:
        assignment.seed = payload.seed
    audit(session, user, "ASSIGNMENT_PUBLISHED", "assignment", assignment.id, request_id(request), {"seed": assignment.seed})
    queue_assignment_notifications(
        session,
        assignment,
        "ASSIGNMENT_PUBLISHED",
        "มีงานประเมินใหม่ กรุณาเปิด PairEval เพื่อตรวจรายละเอียดและดำเนินการ",
    )
    session.commit()
    return {
        "assignmentId": assignment.id,
        "status": assignment.status,
        "pairAssignments": pair_count,
        "feasibility": feasibility_for_assignment(session, assignment),
    }


@app.get("/api/assignments/{assignment_id}/my-evaluations")
def get_my_evaluations(
    assignment_id: str,
    session: SessionDep,
    user: UserDep,
    side: Literal["GROUP", "INDIVIDUAL"] = "GROUP",
):
    assignment, classroom, membership = assignment_access(
        session, assignment_id, user, {"OWNER", "INSTRUCTOR", "STUDENT"}
    )
    if assignment.status not in {"PUBLISHED", "CLOSED", "FINALIZED"}:
        raise ApiError(409, "ASSIGNMENT_NOT_PUBLISHED", "This assignment is not open for evaluation.")
    criteria = {
        item.id: item
        for item in session.scalars(select(CriterionRecord).where(CriterionRecord.assignment_id == assignment.id)).all()
    }
    pairs = session.scalars(
        select(PairAssignmentRecord)
        .join(CriterionRecord, CriterionRecord.id == PairAssignmentRecord.criterion_id)
        .where(
            PairAssignmentRecord.assignment_id == assignment.id,
            PairAssignmentRecord.evaluator_id == user.id,
            CriterionRecord.side == side,
        )
        .order_by(PairAssignmentRecord.criterion_id, PairAssignmentRecord.id)
    ).all()
    item_ids = {pair.item_a_id for pair in pairs} | {pair.item_b_id for pair in pairs}
    if side == "GROUP":
        item_records = {
            item.id: {"name": item.name, "artifactUrl": item.artifact_url}
            for item in session.scalars(select(GroupRecord).where(GroupRecord.id.in_(item_ids))).all()
        } if item_ids else {}
    else:
        item_records = {
            item.id: {"name": item.display_name, "artifactUrl": None}
            for item in session.scalars(select(UserRecord).where(UserRecord.id.in_(item_ids))).all()
        } if item_ids else {}
    comparisons = {
        item.pair_assignment_id: item
        for item in session.scalars(
            select(ComparisonRecord).where(ComparisonRecord.pair_assignment_id.in_([pair.id for pair in pairs]))
        ).all()
    } if pairs else {}
    result = []
    for pair in pairs:
        criterion = criteria[pair.criterion_id]
        left_id = pair.display_left_item_id
        right_id = pair.item_b_id if left_id == pair.item_a_id else pair.item_a_id
        current = comparisons.get(pair.id)
        result.append(
            {
                "id": pair.id,
                "criterionId": criterion.id,
                "criterion": criterion.name,
                "prompt": criterion.prompt,
                "left": {"id": left_id, **item_records[left_id]},
                "right": {"id": right_id, **item_records[right_id]},
                "choice": current.choice if current else None,
                "status": current.status if current else "UNANSWERED",
            }
        )
    return {
        "assignment": assignment_payload(session, assignment, user.id),
        "classroom": {"id": classroom.id, "name": classroom.name},
        "groupId": membership.group_id,
        "side": side,
        "readOnly": assignment.status != "PUBLISHED" or evaluation_deadline(assignment, side) <= datetime.now(timezone.utc),
        "disabledReason": (
            "กลุ่มที่มีสมาชิกไม่เกิน 2 คนไม่สามารถประเมินรายบุคคลได้"
            if side == "INDIVIDUAL" and not pairs
            else None
        ),
        "pairs": result,
    }


@app.put("/api/comparisons/{pair_assignment_id}")
def save_comparison(pair_assignment_id: str, payload: ComparisonSave, session: SessionDep, user: UserDep):
    row = session.execute(
        select(
            PairAssignmentRecord,
            AssignmentRecord,
            MembershipRecord,
            CriterionRecord.side,
            ComparisonRecord,
        )
        .join(AssignmentRecord, AssignmentRecord.id == PairAssignmentRecord.assignment_id)
        .join(
            MembershipRecord,
            (MembershipRecord.classroom_id == AssignmentRecord.classroom_id)
            & (MembershipRecord.user_id == user.id),
        )
        .join(CriterionRecord, CriterionRecord.id == PairAssignmentRecord.criterion_id)
        .outerjoin(ComparisonRecord, ComparisonRecord.pair_assignment_id == PairAssignmentRecord.id)
        .where(PairAssignmentRecord.id == pair_assignment_id)
    ).one_or_none()
    if row is None:
        raise ApiError(404, "PAIR_ASSIGNMENT_NOT_FOUND", "Pair assignment was not found.")
    pair, assignment, membership, side, current = row
    if membership.role not in {"OWNER", "INSTRUCTOR", "STUDENT"}:
        raise ApiError(403, "CLASSROOM_FORBIDDEN", "Your role cannot perform this action.")
    if pair.evaluator_id != user.id:
        raise ApiError(403, "PAIR_ASSIGNMENT_FORBIDDEN", "This pair belongs to another evaluator.")
    if assignment.status != "PUBLISHED" or evaluation_deadline(assignment, side or "GROUP") <= datetime.now(timezone.utc):
        raise ApiError(409, "EVALUATION_CLOSED", "The evaluation deadline has passed.")
    now = datetime.now(timezone.utc)
    if current is None:
        current = ComparisonRecord(
            pair_assignment_id=pair.id,
            evaluator_id=user.id,
            choice=payload.choice,
            status="DRAFT",
            time_on_task_ms=payload.time_on_task_ms,
            saved_at=now,
        )
        session.add(current)
    else:
        current.choice = payload.choice
        current.status = "DRAFT"
        current.time_on_task_ms = payload.time_on_task_ms
        current.saved_at = now
        current.submitted_at = None
    session.commit()
    return {"pairAssignmentId": pair.id, "choice": current.choice, "status": current.status, "savedAt": aware_utc(current.saved_at).isoformat()}


@app.post("/api/assignments/{assignment_id}/submissions", status_code=201)
def submit_evaluation(
    assignment_id: str,
    payload: SubmissionRequest,
    session: SessionDep,
    user: UserDep,
    idempotency_key: Annotated[str, Header(alias="Idempotency-Key", min_length=8)],
):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR", "STUDENT"})
    existing = session.scalar(
        select(SubmissionRevisionRecord).where(
            SubmissionRevisionRecord.assignment_id == assignment.id,
            SubmissionRevisionRecord.evaluator_id == user.id,
            SubmissionRevisionRecord.idempotency_key == idempotency_key,
        )
    )
    if existing:
        if existing.side != payload.side:
            raise ApiError(409, "IDEMPOTENCY_KEY_REUSED", "This Idempotency-Key was already used for another side.")
        return {
            "assignmentId": assignment.id,
            "side": existing.side,
            "revision": existing.revision,
            "submittedCount": existing.submitted_count,
            "unansweredCount": existing.unanswered_count,
            "submittedAt": aware_utc(existing.submitted_at).isoformat(),
        }
    if assignment.status != "PUBLISHED" or evaluation_deadline(assignment, payload.side) <= datetime.now(timezone.utc):
        raise ApiError(409, "EVALUATION_CLOSED", "The evaluation deadline has passed.")
    pairs = session.scalars(
        select(PairAssignmentRecord).where(
            PairAssignmentRecord.assignment_id == assignment.id,
            PairAssignmentRecord.evaluator_id == user.id,
            PairAssignmentRecord.criterion_id.in_(
                select(CriterionRecord.id).where(
                    CriterionRecord.assignment_id == assignment.id,
                    CriterionRecord.side == payload.side,
                )
            ),
        )
    ).all()
    pair_ids = [pair.id for pair in pairs]
    comparisons = session.scalars(
        select(ComparisonRecord).where(ComparisonRecord.pair_assignment_id.in_(pair_ids))
    ).all() if pair_ids else []
    now = datetime.now(timezone.utc)
    answers: dict[str, int] = {}
    time_on_task: dict[str, int] = {}
    for comparison in comparisons:
        comparison.status = "SUBMITTED"
        comparison.submitted_at = now
        answers[comparison.pair_assignment_id] = comparison.choice
        time_on_task[comparison.pair_assignment_id] = comparison.time_on_task_ms
    last_revision = session.scalar(
        select(func.max(SubmissionRevisionRecord.revision)).where(
            SubmissionRevisionRecord.assignment_id == assignment.id,
            SubmissionRevisionRecord.evaluator_id == user.id,
            SubmissionRevisionRecord.side == payload.side,
        )
    ) or 0
    revision = SubmissionRevisionRecord(
        id=str(uuid4()),
        assignment_id=assignment.id,
        evaluator_id=user.id,
        side=payload.side,
        revision=last_revision + 1,
        idempotency_key=idempotency_key,
        submitted_count=len(answers),
        unanswered_count=len(pairs) - len(answers),
        answers_json=json.dumps(answers, sort_keys=True),
        metadata_json=json.dumps({"timeOnTaskMs": time_on_task}, sort_keys=True),
        submitted_at=now,
    )
    session.add(revision)
    session.commit()
    return {
        "assignmentId": assignment.id,
        "side": payload.side,
        "revision": revision.revision,
        "submittedCount": revision.submitted_count,
        "unansweredCount": revision.unanswered_count,
        "submittedAt": now.isoformat(),
    }


@app.get("/api/assignments/{assignment_id}/my-score")
def get_my_score(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, membership = assignment_access(session, assignment_id, user, {"STUDENT"})
    if membership.group_id is None:
        raise ApiError(409, "GROUP_REQUIRED", "The student is not assigned to a group.")
    service = ScoreReportingService(session, assignment)
    row = next(item for item in service.individual_report() if item["itemId"] == user.id)
    individual_scores = service.item_scores("INDIVIDUAL") if assignment.individual_max_score > 0 else []
    own_individual = next((item for item in individual_scores if item["itemId"] == user.id), None)
    individual_private = bool(
        assignment.individual_max_score > 0
        and own_individual
        and any(
            criterion["comparisonCount"] < assignment.min_comparisons
            for criterion in own_individual["criteria"]
        )
    )
    comparison_counts = [
        criterion["comparisonCount"]
        for criterion in (own_individual or {}).get("criteria", [])
    ]
    if individual_private:
        return {
            "state": "INSUFFICIENT_DATA",
            "label": "ยังมีข้อมูลไม่พอ",
            "groupComponent": row["groupComponent"],
            "individualComponent": None,
            "participationRatio": row["participationRatio"],
            "participationMultiplier": row["participationMultiplier"],
            "total": None,
            "comparisonCount": min(comparison_counts, default=0),
            "flags": sorted(set(row["flags"]) | {"ANONYMITY_THRESHOLD"}),
        }
    has_score = row["total"] is not None
    is_final = assignment.status == "FINALIZED" and has_score
    return {
        "state": "FINAL" if is_final else ("INTERIM" if has_score else "INSUFFICIENT_DATA"),
        "label": "คะแนนสิ้นสุด" if is_final else ("ชั่วคราว — อาจเปลี่ยนแปลงได้" if has_score else "ยังมีข้อมูลไม่พอ"),
        "groupComponent": row["groupComponent"],
        "individualComponent": row["individualComponent"],
        "participationRatio": row["participationRatio"],
        "participationMultiplier": row["participationMultiplier"],
        "total": row["total"],
        "comparisonCount": min(comparison_counts, default=0),
        "flags": row["flags"],
    }


@app.get("/api/assignments/{assignment_id}/reports/group")
def group_report(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    return {"assignmentId": assignment.id, "rows": ScoreReportingService(session, assignment).item_scores("GROUP")}


@app.get("/api/assignments/{assignment_id}/reports/individual")
def individual_report(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    return {"assignmentId": assignment.id, "rows": ScoreReportingService(session, assignment).individual_report()}


@app.get("/api/assignments/{assignment_id}/reports/coverage")
def coverage_report(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    return {"assignmentId": assignment.id, "rows": ScoreReportingService(session, assignment).coverage_report()}


@app.get("/api/assignments/{assignment_id}/reports/quality")
def quality_report(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    service = ScoreReportingService(session, assignment)
    flags = [
        {"signal": "LOW_COVERAGE", "side": row["side"], "criterion": row["criterion"], "itemAId": row["itemAId"], "itemBId": row["itemBId"], "action": "ส่งประเมินเพิ่ม"}
        for row in service.coverage_report()
        if row["flags"]
    ]
    flags.extend(QualitySignalService(session, assignment).signals())
    return {"assignmentId": assignment.id, "flags": flags}


def csv_rows_for_report(service: ScoreReportingService, report: str) -> tuple[list[str], list[list[object]]]:
    if report == "GROUP":
        headers = ["group_id", "group_name", "criterion", "quality_index", "comparison_count", "weighted_score", "component", "flags"]
        rows = [
            [item["itemId"], item["itemName"], criterion["criterion"], criterion["qualityIndex"], criterion["comparisonCount"], criterion["weightedScore"], item["component"], "|".join(item["flags"])]
            for item in service.item_scores("GROUP")
            for criterion in item["criteria"]
        ]
    elif report == "INDIVIDUAL":
        headers = ["student_id", "display_name", "group_name", "criterion", "quality_index", "comparison_count", "individual_component", "group_component", "participation", "multiplier", "total", "flags"]
        rows = [
            [item["studentId"], item["itemName"], item["groupName"], criterion.get("criterion", ""), criterion.get("qualityIndex"), criterion.get("comparisonCount", 0), item["individualComponent"], item["groupComponent"], item["participationRatio"], item["participationMultiplier"], item["total"], "|".join(item["flags"])]
            for item in service.individual_report()
            for criterion in (item["criteria"] or [{}])
        ]
    else:
        headers = ["side", "criterion", "item_a_id", "item_b_id", "assigned_coverage", "submitted_coverage", "flags"]
        rows = [
            [item["side"], item["criterion"], item["itemAId"], item["itemBId"], item["assignedCoverage"], item["submittedCoverage"], "|".join(item["flags"])]
            for item in service.coverage_report()
        ]
    return headers, rows


def raw_comparison_rows(
    session: Session,
    assignment: AssignmentRecord,
    include_identities: bool,
) -> tuple[list[str], list[list[object]]]:
    revisions = session.scalars(
        select(SubmissionRevisionRecord)
        .where(SubmissionRevisionRecord.assignment_id == assignment.id)
        .order_by(
            SubmissionRevisionRecord.evaluator_id,
            SubmissionRevisionRecord.side,
            SubmissionRevisionRecord.revision.desc(),
        )
    ).all()
    latest: dict[tuple[str, str], SubmissionRevisionRecord] = {}
    for revision in revisions:
        latest.setdefault((revision.evaluator_id, revision.side), revision)
    pairs = {
        pair.id: pair
        for pair in session.scalars(
            select(PairAssignmentRecord).where(PairAssignmentRecord.assignment_id == assignment.id)
        ).all()
    }
    criteria = {
        criterion.id: criterion
        for criterion in session.scalars(
            select(CriterionRecord).where(CriterionRecord.assignment_id == assignment.id)
        ).all()
    }
    pseudonyms = QualitySignalService(session, assignment)
    rows: list[list[object]] = []
    for revision in latest.values():
        metadata = json.loads(revision.metadata_json or "{}")
        time_on_task = metadata.get("timeOnTaskMs", {})
        evaluator = revision.evaluator_id if include_identities else pseudonyms.pseudonym(revision.evaluator_id)
        for pair_id, choice in json.loads(revision.answers_json).items():
            pair = pairs.get(pair_id)
            if pair is None:
                continue
            criterion = criteria[pair.criterion_id]
            rows.append(
                [
                    evaluator,
                    revision.side,
                    criterion.name,
                    pair.id,
                    pair.item_a_id,
                    pair.item_b_id,
                    choice,
                    time_on_task.get(pair.id),
                    revision.revision,
                    aware_utc(revision.submitted_at).isoformat(),
                ]
            )
    return (
        ["evaluator", "side", "criterion", "pair_id", "item_a_id", "item_b_id", "choice", "time_on_task_ms", "revision", "submitted_at_utc"],
        rows,
    )


def spreadsheet_safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def assignment_slug(assignment: AssignmentRecord) -> str:
    value = re.sub(r"[^a-z0-9]+", "-", assignment.name.lower()).strip("-")
    return value or assignment.id[:8]


def xlsx_report(service: ScoreReportingService, assignment: AssignmentRecord) -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for title, report_name in (
        ("Group Summary", "GROUP"),
        ("Individual Summary", "INDIVIDUAL"),
        ("Pair Coverage", "COVERAGE"),
    ):
        sheet = workbook.create_sheet(title)
        headers, rows = csv_rows_for_report(service, report_name)
        sheet.append([spreadsheet_safe(value) for value in headers])
        for row in rows:
            sheet.append([spreadsheet_safe(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    metadata = workbook.create_sheet("Metadata")
    metadata_rows = [
        ("assignment_id", assignment.id),
        ("assignment_name", assignment.name),
        ("formula_version", assignment.scoring_formula_version),
        ("score_floor", str(assignment.score_floor)),
        ("score_ceiling", str(assignment.score_ceiling)),
        ("completion_threshold", str(assignment.completion_threshold)),
        ("instructor_weight", str(assignment.instructor_weight)),
        ("group_max_score", str(assignment.group_max_score)),
        ("individual_max_score", str(assignment.individual_max_score)),
        ("exported_at_utc", datetime.now(timezone.utc).isoformat()),
    ]
    metadata.append(["key", "value"])
    for row in metadata_rows:
        metadata.append([spreadsheet_safe(value) for value in row])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@app.post("/api/assignments/{assignment_id}/exports")
def export_report(
    assignment_id: str,
    payload: ExportRequest,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    assignment, classroom, membership = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    if payload.include_identities and (
        membership.role != "OWNER" or payload.identity_confirmation != "EXPORT_IDENTITIES"
    ):
        raise ApiError(
            403,
            "IDENTITY_EXPORT_CONFIRMATION_REQUIRED",
            "Only the classroom owner can export identities after explicit confirmation.",
        )
    if payload.include_identities and payload.report != "RAW":
        raise ApiError(422, "IDENTITIES_REQUIRE_RAW_REPORT", "Evaluator identities are available only in the raw report.")
    if payload.format == "XLSX" and payload.report == "RAW":
        raise ApiError(422, "RAW_REPORT_IS_CSV_ONLY", "The raw comparison report is available as CSV only.")
    service = ScoreReportingService(session, assignment)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    base_filename = f"{classroom.slug}_{assignment_slug(assignment)}"
    if payload.format == "XLSX":
        filename = f"{base_filename}_all_{timestamp}.xlsx"
        return Response(
            content=xlsx_report(service, assignment),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    if payload.report == "RAW":
        headers, rows = raw_comparison_rows(session, assignment, payload.include_identities)
    else:
        headers, rows = csv_rows_for_report(service, payload.report)
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([spreadsheet_safe(value) for value in headers])
    writer.writerows([[spreadsheet_safe(value) for value in row] for row in rows])
    filename = f"{base_filename}_{payload.report.lower()}_{timestamp}.csv"
    if payload.include_identities:
        audit(
            session,
            user,
            "EVALUATOR_IDENTITIES_EXPORTED",
            "assignment",
            assignment.id,
            request_id(request),
            after={"report": payload.report, "format": payload.format},
            reason="Owner explicitly confirmed evaluator identity export",
            ip_address=request.client.host if request.client else None,
        )
        session.commit()
    return Response(
        content=("\ufeff" + output.getvalue()).encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/assignments/{assignment_id}:finalize")
def finalize_assignment(
    assignment_id: str,
    payload: FinalizeRequest,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER"})
    if assignment.status not in {"PUBLISHED", "CLOSED", "FINALIZED"}:
        raise ApiError(409, "ASSIGNMENT_NOT_CLOSED", "Only a published or closed assignment can be finalized.")
    if assignment.status == "FINALIZED":
        existing, _ = finalize_scores(session, assignment, user.id)
        return {
            "assignmentId": assignment.id,
            "snapshotId": existing.id,
            "revision": existing.revision,
            "status": assignment.status,
            "finalizedAt": aware_utc(existing.finalized_at).isoformat(),
        }
    latest_deadline = max(aware_utc(assignment.deadline), evaluation_deadline(assignment, "INDIVIDUAL"))
    if latest_deadline > datetime.now(timezone.utc):
        raise ApiError(409, "ASSIGNMENT_STILL_OPEN", "Finalize is available after both evaluation deadlines.")
    before = {
        "status": assignment.status,
        "finalizedAt": assignment.finalized_at.isoformat() if assignment.finalized_at else None,
    }
    snapshot, _ = finalize_scores(session, assignment, user.id)
    audit(
        session,
        user,
        "ASSIGNMENT_FINALIZED",
        "assignment",
        assignment.id,
        request_id(request),
        before=before,
        after={"status": "FINALIZED", "snapshotId": snapshot.id, "revision": snapshot.revision},
        reason=payload.reason,
        ip_address=request.client.host if request.client else None,
        assignment_id=assignment.id,
    )
    queue_assignment_notifications(
        session,
        assignment,
        "SCORES_FINALIZED",
        "งานประเมินได้รับการ Finalize แล้ว กรุณาเปิด PairEval เพื่อดูผลของคุณ",
    )
    session.commit()
    return {
        "assignmentId": assignment.id,
        "snapshotId": snapshot.id,
        "revision": snapshot.revision,
        "status": assignment.status,
        "finalizedAt": aware_utc(snapshot.finalized_at).isoformat(),
    }


@app.post("/api/assignments/{assignment_id}:reopen")
def reopen_assignment(
    assignment_id: str,
    payload: ReasonRequest,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER"})
    if assignment.status != "FINALIZED":
        raise ApiError(409, "ASSIGNMENT_NOT_FINALIZED", "Only a finalized assignment can be reopened.")
    assignment.status = "CLOSED"
    audit(
        session,
        user,
        "ASSIGNMENT_REOPENED",
        "assignment",
        assignment.id,
        request_id(request),
        before={"status": "FINALIZED"},
        after={"status": "CLOSED"},
        reason=payload.reason,
        ip_address=request.client.host if request.client else None,
        assignment_id=assignment.id,
    )
    session.commit()
    return {"assignmentId": assignment.id, "status": assignment.status}


@app.post("/api/assignments/{assignment_id}/overrides", status_code=201)
def create_score_override(
    assignment_id: str,
    payload: ScoreOverrideCreate,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    assignment, classroom, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    maximum = assignment.group_max_score if payload.side == "GROUP" else assignment.individual_max_score
    if payload.override_value > maximum:
        raise ApiError(422, "OVERRIDE_OUT_OF_RANGE", f"Override cannot exceed {maximum}.", "overrideValue")
    if payload.side == "GROUP":
        item_exists = session.scalar(
            select(func.count()).select_from(GroupRecord).where(
                GroupRecord.id == payload.item_id,
                GroupRecord.classroom_id == classroom.id,
            )
        )
    else:
        item_exists = session.scalar(
            select(func.count()).select_from(MembershipRecord).where(
                MembershipRecord.user_id == payload.item_id,
                MembershipRecord.classroom_id == classroom.id,
                MembershipRecord.role == "STUDENT",
            )
        )
    if not item_exists:
        raise ApiError(404, "SCORE_ITEM_NOT_FOUND", "The score item was not found in this classroom.")
    current = next(
        (item for item in ScoreReportingService(session, assignment).item_scores(payload.side) if item["itemId"] == payload.item_id),
        None,
    )
    override = ScoreOverrideRecord(
        id=str(uuid4()),
        assignment_id=assignment.id,
        side=payload.side,
        item_id=payload.item_id,
        original_value=Decimal(str(current["component"])) if current and current["component"] is not None else None,
        override_value=payload.override_value,
        reason=payload.reason,
        created_by=user.id,
    )
    session.add(override)
    audit(
        session,
        user,
        "SCORE_OVERRIDDEN",
        "score",
        payload.item_id,
        request_id(request),
        before={"value": str(override.original_value) if override.original_value is not None else None},
        after={"value": str(payload.override_value), "side": payload.side},
        reason=payload.reason,
        ip_address=request.client.host if request.client else None,
        assignment_id=assignment.id,
    )
    session.commit()
    return {
        "id": override.id,
        "assignmentId": assignment.id,
        "side": override.side,
        "itemId": override.item_id,
        "originalValue": float(override.original_value) if override.original_value is not None else None,
        "overrideValue": float(override.override_value),
        "reason": override.reason,
    }


@app.post("/api/comparisons/{pair_assignment_id}:exclude", status_code=201)
def exclude_comparison(
    pair_assignment_id: str,
    payload: ReasonRequest,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    pair = session.get(PairAssignmentRecord, pair_assignment_id)
    if pair is None:
        raise ApiError(404, "PAIR_ASSIGNMENT_NOT_FOUND", "Pair assignment was not found.")
    assignment, _, _ = assignment_access(session, pair.assignment_id, user, {"OWNER", "INSTRUCTOR"})
    if session.get(ComparisonExclusionRecord, pair.id):
        raise ApiError(409, "COMPARISON_ALREADY_EXCLUDED", "This comparison is already excluded.")
    latest_submitted = any(
        pair.id in answers
        for answers in ScoreReportingService(session, assignment).latest_answers(
            session.scalar(select(CriterionRecord.side).where(CriterionRecord.id == pair.criterion_id)) or "GROUP"
        ).values()
    )
    if not latest_submitted:
        raise ApiError(409, "COMPARISON_NOT_SUBMITTED", "Only a submitted comparison can be excluded.")
    session.add(
        ComparisonExclusionRecord(
            pair_assignment_id=pair.id,
            assignment_id=assignment.id,
            excluded_by=user.id,
            reason=payload.reason,
        )
    )
    audit(
        session,
        user,
        "COMPARISON_EXCLUDED",
        "comparison",
        pair.id,
        request_id(request),
        after={"excluded": True},
        reason=payload.reason,
        ip_address=request.client.host if request.client else None,
        assignment_id=assignment.id,
    )
    session.commit()
    return {"pairAssignmentId": pair.id, "excluded": True, "reason": payload.reason}


@app.post("/api/appeals", status_code=201)
def create_appeal(payload: AppealCreate, request: Request, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, payload.assignment_id, user, {"STUDENT"})
    if assignment.status != "FINALIZED" or assignment.finalized_at is None:
        raise ApiError(409, "APPEAL_NOT_OPEN", "Appeals open only after scores are finalized.")
    if aware_utc(assignment.finalized_at) + timedelta(days=7) < datetime.now(timezone.utc):
        raise ApiError(409, "APPEAL_DEADLINE_PASSED", "The seven-day appeal window has closed.")
    existing = session.scalar(
        select(AppealRecord).where(
            AppealRecord.assignment_id == assignment.id,
            AppealRecord.student_user_id == user.id,
            AppealRecord.status == "OPEN",
        )
    )
    if existing:
        raise ApiError(409, "APPEAL_ALREADY_OPEN", "You already have an open appeal for this assignment.")
    appeal = AppealRecord(
        id=str(uuid4()),
        assignment_id=assignment.id,
        student_user_id=user.id,
        message=payload.message.strip(),
    )
    session.add(appeal)
    audit(
        session,
        user,
        "APPEAL_CREATED",
        "appeal",
        appeal.id,
        request_id(request),
        after={"status": "OPEN"},
        assignment_id=assignment.id,
    )
    session.commit()
    return {"id": appeal.id, "assignmentId": assignment.id, "status": appeal.status, "message": appeal.message}


@app.get("/api/assignments/{assignment_id}/appeals")
def list_appeals(assignment_id: str, session: SessionDep, user: UserDep):
    assignment, _, _ = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    appeals = session.scalars(
        select(AppealRecord).where(AppealRecord.assignment_id == assignment.id).order_by(AppealRecord.created_at.desc())
    ).all()
    users = {
        item.id: item
        for item in session.scalars(select(UserRecord).where(UserRecord.id.in_([appeal.student_user_id for appeal in appeals]))).all()
    } if appeals else {}
    return [
        {
            "id": appeal.id,
            "studentId": appeal.student_user_id,
            "studentName": users[appeal.student_user_id].display_name,
            "message": appeal.message,
            "status": appeal.status,
            "resolution": appeal.resolution,
            "createdAt": aware_utc(appeal.created_at).isoformat(),
        }
        for appeal in appeals
    ]


@app.post("/api/appeals/{appeal_id}:resolve")
def resolve_appeal(
    appeal_id: str,
    payload: AppealResolve,
    request: Request,
    session: SessionDep,
    user: UserDep,
):
    appeal = session.get(AppealRecord, appeal_id)
    if appeal is None:
        raise ApiError(404, "APPEAL_NOT_FOUND", "Appeal was not found.")
    assignment, _, _ = assignment_access(session, appeal.assignment_id, user, {"OWNER", "INSTRUCTOR"})
    if appeal.status != "OPEN":
        raise ApiError(409, "APPEAL_ALREADY_RESOLVED", "This appeal has already been resolved.")
    appeal.status = payload.status
    appeal.resolution = payload.resolution.strip()
    appeal.resolved_by = user.id
    appeal.resolved_at = datetime.now(timezone.utc)
    audit(
        session,
        user,
        "APPEAL_RESOLVED",
        "appeal",
        appeal.id,
        request_id(request),
        before={"status": "OPEN"},
        after={"status": appeal.status},
        reason=appeal.resolution,
        ip_address=request.client.host if request.client else None,
        assignment_id=assignment.id,
    )
    session.commit()
    return {"id": appeal.id, "assignmentId": assignment.id, "status": appeal.status, "resolution": appeal.resolution}


@app.get("/api/assignments/{assignment_id}/audit")
def list_audit_events(assignment_id: str, request: Request, session: SessionDep, user: UserDep):
    assignment, _, membership = assignment_access(session, assignment_id, user, {"OWNER", "INSTRUCTOR"})
    events = session.scalars(
        select(AuditRecord)
        .where(AuditRecord.assignment_id == assignment.id)
        .order_by(AuditRecord.created_at.desc())
    ).all()
    if membership.role == "OWNER":
        audit(
            session,
            user,
            "AUDIT_IDENTITY_ACCESSED",
            "assignment",
            assignment.id,
            request_id(request),
            reason="Owner viewed audit actors",
            ip_address=request.client.host if request.client else None,
        )
        session.commit()
    return [
        {
            "id": event.id,
            "actor": event.actor_user_id if membership.role == "OWNER" else QualitySignalService(session, assignment).pseudonym(event.actor_user_id),
            "action": event.action,
            "resourceType": event.resource_type,
            "resourceId": event.resource_id,
            "before": json.loads(event.before_json or "{}"),
            "after": json.loads(event.after_json or "{}"),
            "reason": event.reason,
            "occurredAt": aware_utc(event.created_at).isoformat(),
        }
        for event in events
    ]


@app.get("/api/privacy/notice")
def privacy_notice():
    return {
        "version": "v1.0",
        "summary": "PairEval stores university identity, evaluations and time-on-task for academic assessment.",
        "retention": {"identityAndEvaluations": "2 academic years", "timeOnTask": "1 academic year"},
        "visibility": "Students see only their aggregate result; instructors see reports under classroom authorization.",
    }


@app.post("/api/privacy/acknowledgements", status_code=201)
def acknowledge_privacy(payload: PrivacyAcknowledgement, session: SessionDep, user: UserDep):
    acknowledgement = session.get(PrivacyAcknowledgementRecord, user.id)
    if acknowledgement is None:
        acknowledgement = PrivacyAcknowledgementRecord(user_id=user.id, notice_version=payload.notice_version)
        session.add(acknowledgement)
        session.commit()
    return {"noticeVersion": acknowledgement.notice_version, "acknowledgedAt": aware_utc(acknowledgement.acknowledged_at).isoformat()}


@app.get("/api/me/data-export")
def export_my_data(session: SessionDep, user: UserDep):
    memberships = session.scalars(select(MembershipRecord).where(MembershipRecord.user_id == user.id)).all()
    submissions = session.scalars(
        select(SubmissionRevisionRecord).where(SubmissionRevisionRecord.evaluator_id == user.id).order_by(SubmissionRevisionRecord.submitted_at)
    ).all()
    appeals = session.scalars(select(AppealRecord).where(AppealRecord.student_user_id == user.id)).all()
    return {
        "user": {"id": user.id, "email": user.email, "displayName": user.display_name},
        "memberships": [
            {"classroomId": item.classroom_id, "role": item.role, "groupId": item.group_id, "studentId": item.student_id, "status": item.status}
            for item in memberships
        ],
        "submissions": [
            {"assignmentId": item.assignment_id, "side": item.side, "revision": item.revision, "answers": json.loads(item.answers_json), "submittedAt": aware_utc(item.submitted_at).isoformat()}
            for item in submissions
        ],
        "appeals": [
            {"assignmentId": item.assignment_id, "message": item.message, "status": item.status, "resolution": item.resolution}
            for item in appeals
        ],
    }


@app.get("/api/notifications")
def list_notifications(session: SessionDep, user: UserDep):
    notifications = session.scalars(
        select(NotificationRecord)
        .where(NotificationRecord.user_id == user.id)
        .order_by(NotificationRecord.created_at.desc())
        .limit(50)
    ).all()
    return [
        {
            "id": item.id,
            "type": item.notification_type,
            **json.loads(item.payload_json),
            "readAt": aware_utc(item.read_at).isoformat() if item.read_at else None,
            "createdAt": aware_utc(item.created_at).isoformat(),
        }
        for item in notifications
    ]


@app.post("/api/maintenance/run")
def run_owner_maintenance(session: SessionDep, user: UserDep):
    classroom_ids = list(
        session.scalars(
            select(MembershipRecord.classroom_id).where(
                MembershipRecord.user_id == user.id,
                MembershipRecord.role == "OWNER",
            )
        ).all()
    )
    if not classroom_ids:
        raise ApiError(403, "OWNER_REQUIRED", "Only a classroom owner can run maintenance.")
    return run_maintenance(session, classroom_ids=classroom_ids)


@app.post("/api/notifications/{notification_id}:read")
def mark_notification_read(
    notification_id: str,
    payload: NotificationRead,
    session: SessionDep,
    user: UserDep,
):
    notification = session.get(NotificationRecord, notification_id)
    if notification is None or notification.user_id != user.id:
        raise ApiError(404, "NOTIFICATION_NOT_FOUND", "Notification was not found.")
    notification.read_at = datetime.now(timezone.utc) if payload.read else None
    session.commit()
    return {
        "id": notification.id,
        "readAt": aware_utc(notification.read_at).isoformat() if notification.read_at else None,
    }


@app.get("/api/instructor/classrooms/{classroom_id}/summary")
def instructor_summary(classroom_id: str, session: SessionDep, user: UserDep):
    classroom, _ = classroom_membership(session, classroom_id, user, {"OWNER", "INSTRUCTOR"})
    students = session.scalar(
        select(func.count()).select_from(MembershipRecord).where(
            MembershipRecord.classroom_id == classroom.id,
            MembershipRecord.role == "STUDENT",
        )
    ) or 0
    groups = session.scalar(select(func.count()).select_from(GroupRecord).where(GroupRecord.classroom_id == classroom.id)) or 0
    assignments = session.scalars(
        select(AssignmentRecord).where(AssignmentRecord.classroom_id == classroom.id).order_by(AssignmentRecord.created_at.desc())
    ).all()
    return {
        "classroom": {
            "id": classroom.id,
            "name": classroom.name,
            "slug": classroom.slug,
            "timezone": classroom.timezone,
        },
        "studentCount": students,
        "groupCount": groups,
        "assignments": [assignment_payload(session, assignment, user.id) for assignment in assignments],
    }
