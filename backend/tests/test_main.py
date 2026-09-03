import io
from datetime import datetime, timedelta, timezone
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import DatabaseError
import pytest

from main import app
from src.persistence.database import SessionLocal
from src.persistence.models import (
    AssignmentRecord,
    AuditRecord,
    ComparisonRecord,
    MembershipRecord,
    NotificationRecord,
    ScoreSnapshotRecord,
    SubmissionRevisionRecord,
    UserRecord,
)
from src.persistence.seed import stable_id
from src.services.maintenance_service import run_maintenance


client = TestClient(app)
INSTRUCTOR = {"X-User-Id": "instructor-demo"}
STUDENT = {"X-User-Id": "student-09"}


def test_root_and_health_describe_local_m4_service():
    response = client.get("/")
    assert response.status_code == 200
    assert response.json() == {
        "service": "PairEval API",
        "version": "2.0.0-m4-local",
        "status": "ok",
        "mode": "demo",
    }
    assert response.headers["X-Request-ID"]

    health = client.get("/health")
    assert health.json() == {"status": "ok", "database": "connected"}


def test_feasibility_exposes_reduced_coverage_reason_to_instructor():
    response = client.get("/api/assignments/demo-assignment/feasibility", headers=INSTRUCTOR)
    assert response.status_code == 200
    body = response.json()
    assert body["targetCoverage"] == 5
    assert body["actualCoverage"] == 4
    assert body["workload"] == 1
    assert body["reduced"] is True


def test_unknown_assignment_uses_stable_error_envelope():
    response = client.get("/api/assignments/missing/feasibility", headers=INSTRUCTOR)
    assert response.status_code == 404
    body = response.json()["error"]
    assert body["code"] == "ASSIGNMENT_NOT_FOUND"
    assert body["requestId"]


def test_request_validation_uses_stable_error_envelope():
    response = client.post(
        "/api/classrooms",
        headers=INSTRUCTOR,
        json={"name": "Valid name", "slug": "INVALID SLUG", "timezone": "Not/A-Timezone"},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "REQUEST_VALIDATION_FAILED"
    assert response.json()["error"]["requestId"]
    assert {item["field"] for item in response.json()["error"]["details"]} == {"slug", "timezone"}


def test_student_cannot_call_instructor_feasibility_endpoint():
    response = client.get("/api/assignments/demo-assignment/feasibility", headers=STUDENT)
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "CLASSROOM_FORBIDDEN"


def test_student_draft_submit_and_score_payload_never_exposes_evaluator_identity():
    evaluation = client.get("/api/assignments/demo-assignment/my-evaluations", headers=STUDENT)
    assert evaluation.status_code == 200
    pairs = evaluation.json()["pairs"]
    assert pairs
    pair_id = pairs[0]["id"]

    saved = client.put(
        f"/api/comparisons/{pair_id}",
        headers=STUDENT,
        json={"choice": 3, "timeOnTaskMs": 1200},
    )
    assert saved.status_code == 200
    assert saved.json()["status"] == "DRAFT"

    submitted = client.post(
        "/api/assignments/demo-assignment/submissions",
        headers={**STUDENT, "Idempotency-Key": str(uuid4())},
        json={"side": "GROUP"},
    )
    assert submitted.status_code == 201
    assert submitted.json()["revision"] >= 1

    score = client.get("/api/assignments/demo-assignment/my-score", headers=STUDENT)
    assert score.status_code == 200
    assert "evaluator" not in str(score.json()).lower()
    assert score.json()["participationRatio"] > 0


def test_editing_after_submit_keeps_latest_revision_effective_until_resubmit():
    evaluation = client.get("/api/assignments/demo-assignment/my-evaluations", headers=STUDENT).json()
    pair_id = evaluation["pairs"][0]["id"]
    client.put(f"/api/comparisons/{pair_id}", headers=STUDENT, json={"choice": 2})
    first_key = str(uuid4())
    first = client.post(
        "/api/assignments/demo-assignment/submissions",
        headers={**STUDENT, "Idempotency-Key": first_key},
        json={"side": "GROUP"},
    )
    assert first.status_code == 201
    score_before_draft = client.get("/api/assignments/demo-assignment/my-score", headers=STUDENT).json()

    changed = client.put(f"/api/comparisons/{pair_id}", headers=STUDENT, json={"choice": 6})
    assert changed.json()["status"] == "DRAFT"
    score_with_new_draft = client.get("/api/assignments/demo-assignment/my-score", headers=STUDENT).json()
    assert score_with_new_draft == score_before_draft

    with SessionLocal() as session:
        comparison = session.get(ComparisonRecord, pair_id)
        assert comparison is not None and comparison.choice == 6 and comparison.status == "DRAFT"
        revision = session.scalar(
            select(SubmissionRevisionRecord).where(SubmissionRevisionRecord.idempotency_key == first_key)
        )
        assert revision is not None and f'"{pair_id}": 2' in revision.answers_json


def test_successful_submission_retry_returns_original_result_after_deadline():
    evaluation = client.get("/api/assignments/demo-assignment/my-evaluations", headers=STUDENT).json()
    pair_id = evaluation["pairs"][0]["id"]
    client.put(f"/api/comparisons/{pair_id}", headers=STUDENT, json={"choice": 3})
    key = str(uuid4())
    first = client.post(
        "/api/assignments/demo-assignment/submissions",
        headers={**STUDENT, "Idempotency-Key": key},
        json={"side": "GROUP"},
    )
    assert first.status_code == 201

    with SessionLocal() as session:
        assignment = session.get(AssignmentRecord, "demo-assignment")
        assert assignment is not None
        original_deadline = assignment.deadline
        assignment.deadline = datetime.now(timezone.utc) - timedelta(seconds=1)
        session.commit()
    try:
        retried = client.post(
            "/api/assignments/demo-assignment/submissions",
            headers={**STUDENT, "Idempotency-Key": key},
            json={"side": "GROUP"},
        )
        assert retried.status_code == 201
        assert retried.json() == first.json()
    finally:
        with SessionLocal() as session:
            assignment = session.get(AssignmentRecord, "demo-assignment")
            assert assignment is not None
            assignment.deadline = original_deadline
            session.commit()


def test_instructor_completes_classroom_roster_assignment_publish_flow_atomically():
    suffix = uuid4().hex[:8]
    classroom = client.post(
        "/api/classrooms",
        headers=INSTRUCTOR,
        json={"name": f"M1 Test {suffix}", "slug": f"m1-test-{suffix}", "timezone": "Asia/Bangkok"},
    )
    assert classroom.status_code == 201
    classroom_id = classroom.json()["id"]

    csv_text = "email,group_name,student_id,display_name,artifact_url\n" + "\n".join(
        [
            f"a1-{suffix}@uni.example,Alpha,1,A One,https://example.com/a",
            f"a2-{suffix}@uni.example,Alpha,2,A Two,https://example.com/a",
            f"b1-{suffix}@uni.example,Beta,3,B One,https://example.com/b",
            f"b2-{suffix}@uni.example,Beta,4,B Two,https://example.com/b",
            f"c1-{suffix}@uni.example,Gamma,5,C One,https://example.com/c",
            f"c2-{suffix}@uni.example,Gamma,6,C Two,https://example.com/c",
        ]
    )
    roster = client.post(
        f"/api/classrooms/{classroom_id}/roster:import",
        headers=INSTRUCTOR,
        files={"file": ("roster.csv", csv_text, "text/csv")},
    )
    assert roster.status_code == 201
    assert roster.json() == {
        "classroomId": classroom_id,
        "studentCount": 6,
        "groupCount": 3,
        "atomic": True,
    }

    assignment = client.post(
        f"/api/classrooms/{classroom_id}/assignments",
        headers=INSTRUCTOR,
        json={
            "name": "M1 Pairwise Review",
            "description": "Integration test",
            "deadline": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
            "groupMaxScore": 20,
            "criterionName": "UX",
            "criterionPrompt": "ผลงานใดใช้งานได้ชัดเจนกว่า?",
        },
    )
    assert assignment.status_code == 201
    assignment_id = assignment.json()["id"]

    published = client.post(
        f"/api/assignments/{assignment_id}:publish",
        headers=INSTRUCTOR,
        json={},
    )
    assert published.status_code == 201
    assert published.json()["status"] == "PUBLISHED"
    assert published.json()["pairAssignments"] == 6


def test_m2_individual_evaluation_reports_and_csv_export_flow():
    suffix = uuid4().hex[:8]
    classroom = client.post(
        "/api/classrooms",
        headers=INSTRUCTOR,
        json={"name": f"M2 Test {suffix}", "slug": f"m2-test-{suffix}", "timezone": "Asia/Bangkok"},
    )
    assert classroom.status_code == 201
    classroom_id = classroom.json()["id"]
    roster_rows = []
    for group_index, group_name in enumerate(("Alpha", "Beta", "Gamma")):
        for member_index in range(5):
            number = group_index * 5 + member_index + 1
            roster_rows.append(
                f"m2-{suffix}-{number}@uni.example,{group_name},{number},M2 Student {number},https://example.com/{group_name.lower()}"
            )
    roster = client.post(
        f"/api/classrooms/{classroom_id}/roster:import",
        headers=INSTRUCTOR,
        files={
            "file": (
                "roster.csv",
                "email,group_name,student_id,display_name,artifact_url\n" + "\n".join(roster_rows),
                "text/csv",
            )
        },
    )
    assert roster.status_code == 201

    deadline = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()
    assignment = client.post(
        f"/api/classrooms/{classroom_id}/assignments",
        headers=INSTRUCTOR,
        json={
            "name": "M2 Complete Flow",
            "description": "Group and individual evaluation",
            "deadline": deadline,
            "individualDeadline": deadline,
            "groupMaxScore": 15,
            "individualMaxScore": 5,
            "criteria": [
                {"name": "UX", "prompt": "ผลงานกลุ่มใดใช้งานชัดเจนกว่า?", "weightPct": 100, "side": "GROUP"},
                {"name": "Teamwork", "prompt": "สมาชิกคนใดมีส่วนร่วมมากกว่า?", "weightPct": 100, "side": "INDIVIDUAL"},
            ],
        },
    )
    assert assignment.status_code == 201
    assignment_id = assignment.json()["id"]
    published = client.post(f"/api/assignments/{assignment_id}:publish", headers=INSTRUCTOR, json={})
    assert published.status_code == 201
    assert published.json()["pairAssignments"] == 105
    assert published.json()["feasibility"]["individual"][0]["workload"] == 6

    student_id = stable_id("user", f"m2-{suffix}-1@uni.example")
    student_headers = {"X-User-Id": student_id}
    group_bundle = client.get(
        f"/api/assignments/{assignment_id}/my-evaluations?side=GROUP", headers=student_headers
    )
    individual_bundle = client.get(
        f"/api/assignments/{assignment_id}/my-evaluations?side=INDIVIDUAL", headers=student_headers
    )
    assert group_bundle.status_code == 200 and len(group_bundle.json()["pairs"]) == 1
    assert individual_bundle.status_code == 200 and len(individual_bundle.json()["pairs"]) == 6
    assert all(pair["left"]["id"] != student_id and pair["right"]["id"] != student_id for pair in individual_bundle.json()["pairs"])

    for pair in individual_bundle.json()["pairs"]:
        saved = client.put(f"/api/comparisons/{pair['id']}", headers=student_headers, json={"choice": 3})
        assert saved.status_code == 200
    submitted = client.post(
        f"/api/assignments/{assignment_id}/submissions",
        headers={**student_headers, "Idempotency-Key": str(uuid4())},
        json={"side": "INDIVIDUAL"},
    )
    assert submitted.status_code == 201
    assert submitted.json()["submittedCount"] == 6
    assert submitted.json()["unansweredCount"] == 0

    # Complete both sides for the class so the privacy threshold is satisfied
    # and the combined M2 score can be exercised end to end.
    for number in range(1, 16):
        evaluator_id = stable_id("user", f"m2-{suffix}-{number}@uni.example")
        evaluator_headers = {"X-User-Id": evaluator_id}
        for side in ("GROUP", "INDIVIDUAL"):
            bundle = client.get(
                f"/api/assignments/{assignment_id}/my-evaluations?side={side}", headers=evaluator_headers
            ).json()
            for pair in bundle["pairs"]:
                assert client.put(
                    f"/api/comparisons/{pair['id']}", headers=evaluator_headers, json={"choice": 3}
                ).status_code == 200
            assert client.post(
                f"/api/assignments/{assignment_id}/submissions",
                headers={**evaluator_headers, "Idempotency-Key": str(uuid4())},
                json={"side": side},
            ).status_code == 201

    score = client.get(f"/api/assignments/{assignment_id}/my-score", headers=student_headers)
    assert score.status_code == 200
    assert score.json()["state"] == "INTERIM"
    assert score.json()["groupComponent"] is not None
    assert score.json()["individualComponent"] is not None
    assert score.json()["participationRatio"] == 1.0
    assert 12 <= score.json()["total"] <= 20

    group_report = client.get(f"/api/assignments/{assignment_id}/reports/group", headers=INSTRUCTOR)
    individual_report = client.get(f"/api/assignments/{assignment_id}/reports/individual", headers=INSTRUCTOR)
    coverage_report = client.get(f"/api/assignments/{assignment_id}/reports/coverage", headers=INSTRUCTOR)
    assert len(group_report.json()["rows"]) == 3
    assert all(row["component"] is not None for row in group_report.json()["rows"])
    assert len(individual_report.json()["rows"]) == 15
    assert len(coverage_report.json()["rows"]) == 33
    assert client.get(f"/api/assignments/{assignment_id}/reports/group", headers=student_headers).status_code == 403

    exported = client.post(
        f"/api/assignments/{assignment_id}/exports",
        headers=INSTRUCTOR,
        json={"format": "CSV", "report": "INDIVIDUAL"},
    )
    assert exported.status_code == 200
    assert exported.content.startswith(b"\xef\xbb\xbf")
    assert "attachment;" in exported.headers["content-disposition"]
    assert "participation" in exported.text


def test_two_group_class_uses_instructor_evaluation_without_self_group_violation():
    suffix = uuid4().hex[:8]
    classroom = client.post(
        "/api/classrooms",
        headers=INSTRUCTOR,
        json={"name": f"Two Groups {suffix}", "slug": f"two-groups-{suffix}", "timezone": "Asia/Bangkok"},
    ).json()
    csv_text = "email,group_name,student_id,display_name,artifact_url\n" + "\n".join(
        [
            f"two-a1-{suffix}@uni.example,Alpha,1,A One,https://example.com/a",
            f"two-a2-{suffix}@uni.example,Alpha,2,A Two,https://example.com/a",
            f"two-b1-{suffix}@uni.example,Beta,3,B One,https://example.com/b",
            f"two-b2-{suffix}@uni.example,Beta,4,B Two,https://example.com/b",
        ]
    )
    assert client.post(
        f"/api/classrooms/{classroom['id']}/roster:import",
        headers=INSTRUCTOR,
        files={"file": ("roster.csv", csv_text, "text/csv")},
    ).status_code == 201
    assignment = client.post(
        f"/api/classrooms/{classroom['id']}/assignments",
        headers=INSTRUCTOR,
        json={
            "name": "Two-group review",
            "deadline": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
            "groupMaxScore": 20,
            "individualMaxScore": 0,
            "criterionName": "Quality",
            "criterionPrompt": "ผลงานกลุ่มใดมีคุณภาพดีกว่า?",
        },
    ).json()
    feasibility = client.get(f"/api/assignments/{assignment['id']}/feasibility", headers=INSTRUCTOR)
    assert feasibility.status_code == 200
    assert feasibility.json()["workload"] == 0
    assert "ผู้สอน" in feasibility.json()["explanation"]
    published = client.post(f"/api/assignments/{assignment['id']}:publish", headers=INSTRUCTOR, json={})
    assert published.status_code == 201
    assert published.json()["pairAssignments"] == 1
    instructor_bundle = client.get(
        f"/api/assignments/{assignment['id']}/my-evaluations?side=GROUP", headers=INSTRUCTOR
    )
    assert instructor_bundle.status_code == 200
    assert len(instructor_bundle.json()["pairs"]) == 1
    assert set((instructor_bundle.json()["pairs"][0]["left"]["name"], instructor_bundle.json()["pairs"][0]["right"]["name"])) == {"Alpha", "Beta"}


def test_m3_quality_override_exclusion_finalize_appeal_audit_and_privacy_flow():
    suffix = uuid4().hex[:8]
    classroom = client.post(
        "/api/classrooms",
        headers=INSTRUCTOR,
        json={"name": f"M3 Test {suffix}", "slug": f"m3-test-{suffix}", "timezone": "Asia/Bangkok"},
    ).json()
    roster_rows = []
    group_names = ("Alpha", "Beta", "Gamma", "Delta")
    for group_index, group_name in enumerate(group_names):
        for member_index in range(2):
            number = group_index * 2 + member_index + 1
            roster_rows.append(
                f"m3-{suffix}-{number}@uni.example,{group_name},{number},M3 Student {number},https://example.com/{group_name.lower()}"
            )
    imported = client.post(
        f"/api/classrooms/{classroom['id']}/roster:import",
        headers=INSTRUCTOR,
        files={
            "file": (
                "roster.csv",
                "email,group_name,student_id,display_name,artifact_url\n" + "\n".join(roster_rows),
                "text/csv",
            )
        },
    )
    assert imported.status_code == 201

    deadline = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()
    assignment = client.post(
        f"/api/classrooms/{classroom['id']}/assignments",
        headers=INSTRUCTOR,
        json={
            "name": "M3 Governance Flow",
            "deadline": deadline,
            "groupMaxScore": 20,
            "individualMaxScore": 0,
            "criteria": [
                {"name": "Quality", "prompt": "ผลงานกลุ่มใดมีคุณภาพดีกว่า?", "weightPct": 100, "side": "GROUP"}
            ],
        },
    ).json()
    assignment_id = assignment["id"]
    assert client.post(f"/api/assignments/{assignment_id}:publish", headers=INSTRUCTOR, json={}).status_code == 201

    first_student_id = stable_id("user", f"m3-{suffix}-1@uni.example")
    first_student_headers = {"X-User-Id": first_student_id}
    published_notifications = client.get("/api/notifications", headers=first_student_headers)
    assert any(item["type"] == "ASSIGNMENT_PUBLISHED" for item in published_notifications.json())
    first_pair_id = None
    for number in range(1, 9):
        student_id = stable_id("user", f"m3-{suffix}-{number}@uni.example")
        student_headers = {"X-User-Id": student_id}
        bundle = client.get(
            f"/api/assignments/{assignment_id}/my-evaluations?side=GROUP",
            headers=student_headers,
        ).json()
        assert len(bundle["pairs"]) == 3
        for pair in bundle["pairs"]:
            if first_pair_id is None:
                first_pair_id = pair["id"]
            assert client.put(
                f"/api/comparisons/{pair['id']}",
                headers=student_headers,
                json={"choice": 1, "timeOnTaskMs": 1000},
            ).status_code == 200
        assert client.post(
            f"/api/assignments/{assignment_id}/submissions",
            headers={**student_headers, "Idempotency-Key": str(uuid4())},
            json={"side": "GROUP"},
        ).status_code == 201

    quality = client.get(f"/api/assignments/{assignment_id}/reports/quality", headers=INSTRUCTOR)
    assert quality.status_code == 200
    signal_names = {flag["signal"] for flag in quality.json()["flags"]}
    assert {"STRAIGHT_LINING", "POSITION_BIAS", "SPEED_RUN"}.issubset(signal_names)
    assert all("student" not in str(flag.get("rater", "")).lower() for flag in quality.json()["flags"])

    report = client.get(f"/api/assignments/{assignment_id}/reports/group", headers=INSTRUCTOR).json()
    overridden_item_id = report["rows"][0]["itemId"]
    overridden = client.post(
        f"/api/assignments/{assignment_id}/overrides",
        headers=INSTRUCTOR,
        json={"side": "GROUP", "itemId": overridden_item_id, "overrideValue": 18.5, "reason": "Confirmed moderation adjustment"},
    )
    assert overridden.status_code == 201
    overridden_row = next(
        row
        for row in client.get(f"/api/assignments/{assignment_id}/reports/group", headers=INSTRUCTOR).json()["rows"]
        if row["itemId"] == overridden_item_id
    )
    assert overridden_row["component"] == 18.5
    assert "OVERRIDDEN" in overridden_row["flags"]

    assert first_pair_id is not None
    excluded = client.post(
        f"/api/comparisons/{first_pair_id}:exclude",
        headers=INSTRUCTOR,
        json={"reason": "Confirmed anomalous response"},
    )
    assert excluded.status_code == 201
    assert excluded.json()["excluded"] is True

    xlsx = client.post(
        f"/api/assignments/{assignment_id}/exports",
        headers=INSTRUCTOR,
        json={"format": "XLSX", "report": "GROUP"},
    )
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content), read_only=True)
    assert workbook.sheetnames == ["Group Summary", "Individual Summary", "Pair Coverage", "Metadata"]
    assert workbook["Metadata"]["B5"].value == "0.600"
    co_teacher_id = f"co-teacher-{suffix}"
    with SessionLocal() as session:
        session.add(UserRecord(id=co_teacher_id, email=f"co-{suffix}@uni.example", display_name="Co Teacher"))
        session.add(
            MembershipRecord(
                id=str(uuid4()),
                classroom_id=classroom["id"],
                user_id=co_teacher_id,
                role="INSTRUCTOR",
            )
        )
        session.commit()
    raw_pseudonymous = client.post(
        f"/api/assignments/{assignment_id}/exports",
        headers={"X-User-Id": co_teacher_id},
        json={"format": "CSV", "report": "RAW"},
    )
    assert raw_pseudonymous.status_code == 200
    assert "RATER-" in raw_pseudonymous.text
    assert first_student_id not in raw_pseudonymous.text
    assert client.post(
        f"/api/assignments/{assignment_id}/exports",
        headers={"X-User-Id": co_teacher_id},
        json={"format": "CSV", "report": "RAW", "includeIdentities": True, "identityConfirmation": "EXPORT_IDENTITIES"},
    ).status_code == 403
    raw_identified = client.post(
        f"/api/assignments/{assignment_id}/exports",
        headers=INSTRUCTOR,
        json={"format": "CSV", "report": "RAW", "includeIdentities": True, "identityConfirmation": "EXPORT_IDENTITIES"},
    )
    assert raw_identified.status_code == 200
    assert first_student_id in raw_identified.text

    with SessionLocal() as session:
        stored_assignment = session.get(AssignmentRecord, assignment_id)
        assert stored_assignment is not None
        stored_assignment.deadline = datetime.now(timezone.utc) - timedelta(minutes=1)
        stored_assignment.individual_deadline = stored_assignment.deadline
        session.commit()

    forbidden_finalize = client.post(
        f"/api/assignments/{assignment_id}:finalize",
        headers={"X-User-Id": co_teacher_id},
        json={"reason": "Co-teacher must not finalize"},
    )
    assert forbidden_finalize.status_code == 403

    finalized = client.post(
        f"/api/assignments/{assignment_id}:finalize",
        headers=INSTRUCTOR,
        json={"reason": "Evaluation window complete"},
    )
    assert finalized.status_code == 200
    assert finalized.json()["status"] == "FINALIZED"
    assert client.post(
        f"/api/assignments/{assignment_id}:finalize", headers=INSTRUCTOR, json={}
    ).json() == finalized.json()

    student_headers = first_student_headers
    final_score = client.get(f"/api/assignments/{assignment_id}/my-score", headers=student_headers)
    assert final_score.status_code == 200
    assert final_score.json()["state"] == "FINAL"
    assert final_score.json()["label"] == "คะแนนสิ้นสุด"
    finalized_notifications = client.get("/api/notifications", headers=student_headers).json()
    final_notice = next(item for item in finalized_notifications if item["type"] == "SCORES_FINALIZED")
    assert "score" not in final_notice["message"].lower()
    marked_read = client.post(
        f"/api/notifications/{final_notice['id']}:read",
        headers=student_headers,
        json={"read": True},
    )
    assert marked_read.status_code == 200 and marked_read.json()["readAt"]

    appeal = client.post(
        "/api/appeals",
        headers=student_headers,
        json={"assignmentId": assignment_id, "message": "Please review my final score calculation."},
    )
    assert appeal.status_code == 201
    appeal_id = appeal.json()["id"]
    assert client.get(f"/api/assignments/{assignment_id}/appeals", headers=INSTRUCTOR).status_code == 200
    resolved = client.post(
        f"/api/appeals/{appeal_id}:resolve",
        headers=INSTRUCTOR,
        json={"status": "RESOLVED", "resolution": "Reviewed against the immutable score snapshot."},
    )
    assert resolved.status_code == 200
    assert resolved.json()["status"] == "RESOLVED"

    audit_events = client.get(f"/api/assignments/{assignment_id}/audit", headers=INSTRUCTOR)
    assert audit_events.status_code == 200
    audit_actions = {event["action"] for event in audit_events.json()}
    assert {
        "ASSIGNMENT_PUBLISHED",
        "SCORE_OVERRIDDEN",
        "COMPARISON_EXCLUDED",
        "ASSIGNMENT_FINALIZED",
        "APPEAL_CREATED",
        "APPEAL_RESOLVED",
        "EVALUATOR_IDENTITIES_EXPORTED",
    }.issubset(audit_actions)

    notice = client.get("/api/privacy/notice")
    assert notice.status_code == 200
    assert notice.json()["retention"]["timeOnTask"] == "1 academic year"
    acknowledged = client.post(
        "/api/privacy/acknowledgements",
        headers=student_headers,
        json={"noticeVersion": "v1.0"},
    )
    assert acknowledged.status_code == 201
    personal_export = client.get("/api/me/data-export", headers=student_headers)
    assert personal_export.status_code == 200
    assert personal_export.json()["user"]["id"] == first_student_id
    assert any(item["assignmentId"] == assignment_id for item in personal_export.json()["appeals"])


def test_audit_records_are_append_only_at_database_layer():
    with SessionLocal() as session:
        event = session.scalar(select(AuditRecord).order_by(AuditRecord.created_at))
        assert event is not None
        event.action = "MUTATED"
        with pytest.raises(DatabaseError, match="append-only"):
            session.commit()


def test_maintenance_auto_finalizes_fourteen_days_after_deadline_once():
    suffix = uuid4().hex[:8]
    assignment_id = f"maintenance-{suffix}"
    with SessionLocal() as session:
        session.add(
            AssignmentRecord(
                id=assignment_id,
                classroom_id="demo-classroom",
                name="Maintenance finalization",
                description="",
                status="PUBLISHED",
                group_max_score=0,
                individual_max_score=0,
                deadline=datetime.now(timezone.utc) - timedelta(days=15),
                individual_deadline=datetime.now(timezone.utc) - timedelta(days=15),
            )
        )
        session.commit()
        first = run_maintenance(session)
        second = run_maintenance(session)
        assignment = session.get(AssignmentRecord, assignment_id)
        snapshots = session.scalars(
            select(ScoreSnapshotRecord).where(ScoreSnapshotRecord.assignment_id == assignment_id)
        ).all()
        notifications = session.scalars(
            select(NotificationRecord).where(NotificationRecord.dedupe_key.like(f"SCORES_FINALIZED:{assignment_id}:%"))
        ).all()
    assert first["assignmentsFinalized"] == 1
    assert second["assignmentsFinalized"] == 0
    assert assignment is not None and assignment.status == "FINALIZED"
    assert len(snapshots) == 1
    assert notifications


def test_maintenance_applies_confirmed_retention_periods():
    suffix = uuid4().hex[:8]
    user_id = f"retention-user-{suffix}"
    assignment_id = f"retention-assignment-{suffix}"
    submitted_at = datetime.now(timezone.utc) - timedelta(days=731)
    with SessionLocal() as session:
        session.add(UserRecord(id=user_id, email=f"retention-{suffix}@uni.example", display_name="Retention Student"))
        session.add(
            MembershipRecord(
                id=str(uuid4()),
                classroom_id="demo-classroom",
                user_id=user_id,
                role="STUDENT",
                student_id=f"R-{suffix}",
            )
        )
        session.add(
            AssignmentRecord(
                id=assignment_id,
                classroom_id="demo-classroom",
                name="Retention test",
                description="",
                status="FINALIZED",
                group_max_score=0,
                individual_max_score=0,
                deadline=submitted_at,
                finalized_at=submitted_at,
            )
        )
        session.add(
            SubmissionRevisionRecord(
                id=str(uuid4()),
                assignment_id=assignment_id,
                evaluator_id=user_id,
                side="GROUP",
                revision=1,
                idempotency_key=str(uuid4()),
                submitted_count=0,
                unanswered_count=0,
                answers_json="{}",
                metadata_json='{"timeOnTaskMs": {"pair": 1000}}',
                submitted_at=submitted_at,
            )
        )
        session.commit()
        result = run_maintenance(session)
        user = session.get(UserRecord, user_id)
        membership = session.scalar(select(MembershipRecord).where(MembershipRecord.user_id == user_id))
        revision = session.scalar(
            select(SubmissionRevisionRecord).where(SubmissionRevisionRecord.evaluator_id == user_id)
        )
    assert result["timeOnTaskPurged"] >= 1
    assert result["studentIdentitiesAnonymized"] >= 1
    assert user is not None and user.email.endswith("@anonymized.invalid") and user.display_name.startswith("Anonymous")
    assert membership is not None and membership.student_id is None
    assert revision is not None and "timeOnTaskMs" not in revision.metadata_json
